from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "review_1_proposal_design" / "Review_1_Project_Proposal_and_Design.docx"

BLACK = "000000"
BODY = "242424"
MUTED = "5B6573"
ACCENT = "1F4E79"
LIGHT_BLUE = "EAF2F8"
PALE_BLUE = "F5F9FC"
LIGHT_GRAY = "F1F3F5"
GRID = "C9D3DD"


def set_run_font(run, name="Aptos", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="6", val="single"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), val)
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, color=BODY, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=8.5, color=MUTED)


def add_body(doc, text, style="Body Text", italic=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    if italic:
        r.italic = True
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.keep_together = True
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.keep_together = True
    p.add_run(text)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=font_size)
        set_cell_shading(cell, ACCENT)
        set_cell_border(cell, color=ACCENT, size="6")
        if widths:
            cell.width = Inches(widths[i])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            set_cell_text(cell, value, size=font_size)
            set_cell_shading(cell, PALE_BLUE if row_index % 2 == 0 else "FFFFFF")
            set_cell_border(cell)
            if widths:
                cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_stage(doc, title, description):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_BLUE)
    set_cell_border(cell, color=GRID, size="8")
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.5, color=BLACK, bold=True)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(description)
    set_run_font(r2, size=9, color=MUTED)
    arrow = doc.add_paragraph()
    arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    arrow.paragraph_format.space_after = Pt(2)
    ar = arrow.add_run("|")
    set_run_font(ar, size=11, color=ACCENT, bold=True)
    return table


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F8FA")
    set_cell_border(cell, color=GRID, size="6")
    set_cell_margins(cell, top=160, start=180, bottom=160, end=180)
    cell.text = ""
    for idx, line in enumerate(lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        set_run_font(r, name="Consolas", size=8.3, color=BODY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    left, right = table.rows[0].cells
    left.width = Inches(1.25)
    right.width = Inches(5.5)
    set_cell_shading(left, ACCENT)
    set_cell_shading(right, LIGHT_BLUE)
    set_cell_border(left, color=ACCENT, size="6")
    set_cell_border(right, color=GRID, size="6")
    set_cell_text(left, label, bold=True, color="FFFFFF", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(right, text, color=BODY, size=9.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    body = styles["Body Text"]
    body.font.name = "Aptos"
    body._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    body._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    body.font.size = Pt(10.5)
    body.font.color.rgb = RGBColor.from_string(BODY)
    body.paragraph_format.space_after = Pt(6)
    body.paragraph_format.line_spacing = 1.12

    for name, size, before, after in (("Title", 23, 0, 6), ("Heading 1", 15, 14, 6), ("Heading 2", 12.5, 10, 4), ("Heading 3", 11, 8, 3)):
        style = styles[name]
        style.font.name = "Aptos Display" if name in ("Title", "Heading 1") else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if name == "Title":
            ppr = style._element.get_or_add_pPr()
            for child in list(ppr):
                if child.tag == qn("w:pBdr"):
                    ppr.remove(child)

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(10.2)
        style.font.color.rgb = RGBColor.from_string(BODY)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.08


def configure_page(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("MODEL FORGE  |  REVIEW 1")
    set_run_font(r, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    r = fp.add_run("Compiler Design Laboratory  |  Page ")
    set_run_font(r, size=8.5, color=MUTED)
    add_page_number(fp)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(17)
    r = p.add_run("COMPILER DESIGN LABORATORY")
    set_run_font(r, size=10, color=ACCENT, bold=True)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(4)
    title.add_run("ModelForge")

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(22)
    r = sub.add_run("A Lightweight ONNX to C++ Compiler for Standalone Neural Network Inference")
    set_run_font(r, size=13.5, color=MUTED, bold=False)

    panel = doc.add_table(rows=1, cols=1)
    panel.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = panel.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_border(cell, color=GRID, size="8")
    set_cell_margins(cell, top=170, start=210, bottom=170, end=210)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("REVIEW 1  |  PROJECT PROPOSAL AND DESIGN")
    set_run_font(r, size=10.5, color=ACCENT, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run("Problem definition, system architecture, methodology, feasibility, and initial prototype")
    set_run_font(r2, size=10, color=BODY)

    doc.add_paragraph().paragraph_format.space_after = Pt(13)
    meta = add_table(
        doc,
        ["Submission detail", "Value"],
        [
            ("Student name", "____________________________________________"),
            ("Registration number", "____________________________________________"),
            ("Course", "Compiler Design Laboratory"),
            ("Review", "Review 1 - Project Proposal and Design Review"),
            ("Submission date", "____________________________________________"),
        ],
        widths=[2.1, 4.65],
        font_size=9.5,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(26)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("Project status")
    set_run_font(r, size=9, color=MUTED, bold=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run("Phase 1 - Proposal, system design, and prototype boundary")
    set_run_font(r2, size=10.5, color=BODY)

    doc.add_page_break()


def build_document():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    doc.core_properties.title = "ModelForge Review 1 Project Proposal and Design"
    doc.core_properties.subject = "Compiler Design Laboratory project proposal"
    doc.core_properties.author = ""
    doc.core_properties.comments = ""

    add_cover(doc)

    add_heading(doc, "Review 1 Summary", level=1)
    add_body(doc, "ModelForge is proposed as an individual educational compiler that converts a small, deliberately limited subset of ONNX feed-forward neural-network models into readable standalone C++ inference code. The project applies compiler-design ideas to a machine-learning representation: it loads a graph, validates operators and tensor properties, lowers the graph into a custom intermediate representation, applies understandable optimizations, and generates target code.")
    add_callout(doc, "REVIEW 1 OUTCOME", "The design is technically feasible because the first model is a small dense classifier and the compiler boundary is explicit: ONNX-specific parsing ends at the loader, while validation, IR, optimization, code generation, and numerical checking are separate modules.")
    add_body(doc, "This document covers the Phase 1 deliverables requested by the laboratory manual: project definition, background study, requirements, compiler concepts, methodology, architecture, technology selection, feasibility, and an initial prototype plan. The later reviews are separated into their own project folders so implementation evidence can be added without mixing it with the proposal.")

    add_heading(doc, "1 Project Identity", level=1)
    add_table(
        doc,
        ["Item", "Definition"],
        [
            ("Project title", "ModelForge"),
            ("Subtitle", "A Lightweight ONNX to C++ Compiler for Standalone Neural Network Inference"),
            ("Project type", "Individual Compiler Design Laboratory project"),
            ("Implementation language", "C++17"),
            ("Input", "Supported ONNX model (.onnx)"),
            ("Output", "Standalone C++ inference source (.cpp and .h)"),
            ("Primary demonstration", "Small Iris-style feed-forward classifier"),
        ],
        widths=[2.0, 4.75],
    )

    add_heading(doc, "2 Abstract", level=1)
    add_body(doc, "Modern machine-learning models are often exported to portable formats such as ONNX and executed through a general-purpose inference runtime. For a small neural network, that runtime can hide the translation steps that are valuable in a Compiler Design Laboratory. ModelForge addresses this gap by implementing a lightweight educational compiler for a controlled operator subset.")
    add_body(doc, "The compiler will read a supported ONNX graph, extract its inputs, outputs, operators, tensor metadata, and initializers, and perform semantic validation before code generation. The validated graph will be lowered into a simplified ModelForge intermediate representation. Optimization passes such as constant folding, redundant-operation elimination, identity removal, and selected Dense plus ReLU fusion will then be applied. Finally, the compiler will emit readable C++17 code that can be compiled independently for inference.")
    add_body(doc, "Correctness will be evaluated by comparing the generated program's numerical output with the original model executed by ONNX Runtime. The main contribution is not complete ONNX coverage; it is a transparent, modular demonstration of model loading, semantic analysis, intermediate representation, optimization, code generation, and output preservation.")

    add_heading(doc, "3 Problem Statement", level=1)
    add_body(doc, "Small machine-learning models are commonly executed through full inference runtimes. Although this is practical in production, it makes the compilation path difficult to observe and introduces runtime dependencies for simple feed-forward networks. Students therefore need a compact system through which they can inspect a model graph, validate its semantics, transform it into an intermediate representation, optimize it, and generate target code that preserves the model's output.")
    add_callout(doc, "CORE QUESTION", "How can a small ONNX neural-network model be translated into understandable, standalone C++ inference code while demonstrating semantic validation, intermediate representation, optimization, target-code generation, and correctness verification?")

    add_heading(doc, "4 Motivation", level=1)
    add_body(doc, "Traditional compiler projects often use arithmetic expressions, toy languages, or small scripting languages as input. ModelForge applies the same compiler pipeline to a machine-learning graph, creating a direct connection between course concepts and a modern software artifact. A neural-network graph is also a structured computation: operators consume tensors, produce tensors, and carry attributes and constants that must be validated before execution.")
    add_body(doc, "The project is intentionally educational. Readable generated code is treated as a design goal, and optimization reports will explain what changed in the IR. This keeps the system suitable for a technical demonstration and viva rather than turning it into an opaque performance tool.")

    add_heading(doc, "5 Objectives", level=1)
    objectives = [
        "Read a supported neural-network model stored in ONNX format.",
        "Extract graph inputs, outputs, operators, initializers, weights, biases, types, shapes, and relevant attributes.",
        "Detect unsupported operators and invalid tensor properties with compiler-style diagnostics.",
        "Lower the model into a simplified ModelForge intermediate representation independent of ONNX details.",
        "Implement at least two understandable optimization passes and report their effects.",
        "Generate readable standalone C++17 source and header files for inference.",
        "Compile and execute generated code as part of the integrated project.",
        "Compare generated output with the original ONNX model and report numerical differences.",
        "Maintain modular source code, test cases, documentation, and a development log.",
    ]
    for item in objectives:
        add_number(doc, item)

    add_heading(doc, "6 Scope and Boundaries", level=1)
    add_body(doc, "The initial scope is deliberately narrow so that the project can be implemented and defended by one student. The first milestone focuses on small feed-forward networks, especially an Iris classifier with dimensions 4 -> 8 -> 3.")
    add_table(
        doc,
        ["Included in initial scope", "Deferred or excluded initially"],
        [
            ("Gemm and MatMul", "Convolution-heavy architectures"),
            ("Add, Relu, Sigmoid, and Softmax", "YOLO, ResNet, and other large CNN models"),
            ("Static float32 tensor shapes", "Transformers, attention, LSTM, and GRU"),
            ("Small MLP classifiers", "Dynamic control flow and arbitrary dynamic shapes"),
            ("C++17 CPU code generation", "GPU or CUDA code generation"),
            ("Textual or Graphviz support visualization", "Complete ONNX operator coverage"),
        ],
        widths=[3.35, 3.4],
        font_size=9.0,
    )
    add_body(doc, "Possible later additions include Flatten, Tanh, Identity, a larger MNIST MLP, memory reuse, SIMD generation, and richer visualization. These additions will not be allowed to displace the must-have compiler pipeline.")

    add_heading(doc, "7 Background Study", level=1)
    add_heading(doc, "7.1 ONNX as a portable computation graph", level=2)
    add_body(doc, "ONNX describes a model using graph inputs and outputs, operator nodes, initializers, attributes, and tensor types. The format is serialized using Protocol Buffers and defines an operator set with versioned semantics. This makes ONNX a suitable source representation for ModelForge: the loader can translate a graph into compiler-owned structures while keeping the target generator independent of the serialization format. [1]")
    add_heading(doc, "7.2 Runtime execution as a reference", level=2)
    add_body(doc, "ONNX Runtime provides a reference execution path for supported models and offers APIs for several languages, including C and C++. ModelForge will use the original model executed by ONNX Runtime as a comparison oracle during validation. The runtime is not the generated target; it is the baseline against which generated C++ output is checked. [2]")
    add_heading(doc, "7.3 Build and reproducibility", level=2)
    add_body(doc, "CMake is selected as the build-system generator because it can configure C++ projects for different environments and IDEs while keeping the project structure explicit. The generated model source will be built separately from the compiler executable so the demonstration can show both compilation stages. [3]")
    add_callout(doc, "DESIGN CHOICE", "The compiler remains a C++17 project. Python may be used only as an optional helper for exporting or preparing small test models; it is not the implementation language of the compiler.")

    add_heading(doc, "8 Compiler Design Concepts Involved", level=1)
    add_table(
        doc,
        ["Compiler concept", "ModelForge application"],
        [
            ("Source representation", "ONNX graph with nodes, tensors, initializers, and attributes."),
            ("Parsing and model loading", "Read the serialized graph and convert it into C++ model structures."),
            ("Semantic analysis", "Check supported operators, data types, dimensions, required tensors, and attributes."),
            ("Intermediate representation", "Represent inputs, constants, dense operations, activations, Softmax, and return values using ModelForge IR."),
            ("Optimization", "Perform constant folding, redundant-operation elimination, identity removal, and selected operator fusion."),
            ("Target-code generation", "Emit readable C++17 functions, buffers, weights, and activation helpers."),
            ("Execution", "Compile the generated source and run standalone inference."),
            ("Verification", "Compare original ONNX output and generated C++ output within a floating-point tolerance."),
            ("Error handling", "Report unsupported operators, missing tensors, shape mismatches, and invalid attributes clearly."),
        ],
        widths=[2.05, 4.7],
        font_size=8.9,
    )
    add_body(doc, "Lexical analysis and grammar parsing are not the main focus because ONNX is a serialized graph rather than a source language written by the user. The project demonstrates the compiler concepts that fit graph translation most directly, especially semantic analysis, IR construction, optimization, and target generation.")

    add_heading(doc, "9 Proposed Methodology", level=1)
    add_body(doc, "ModelForge follows a staged pipeline. Each stage has a narrow responsibility and passes a well-defined data structure to the next stage.")
    stages = [
        ("1  Load", "Read ONNX graph and metadata"),
        ("2  Validate", "Check supported semantics"),
        ("3  Lower", "Build ModelForge IR"),
        ("4  Optimize", "Transform and report IR"),
        ("5  Generate", "Emit C++17 source"),
        ("6  Verify", "Compare inference output"),
    ]
    for title, desc in stages:
        add_stage(doc, title, desc)
    add_heading(doc, "9.1 Validation strategy", level=2)
    add_body(doc, "Every supported model will be checked before code generation. The validator will reject an operator outside the supported set, incompatible matrix dimensions, unsupported data types, missing initializers, duplicate tensor names, missing outputs, and invalid Softmax axes. Diagnostics will name the failing node or tensor and explain the expected condition.")
    add_heading(doc, "9.2 Optimization strategy", level=2)
    add_body(doc, "Optimization passes will operate on the custom IR rather than directly on ONNX nodes. This provides a stable boundary for reporting before-and-after instruction counts and keeps ONNX parsing separate from compiler transformations. Correctness has priority over any measured speedup; the report will distinguish an IR simplification from a runtime performance claim.")
    add_heading(doc, "9.3 Correctness strategy", level=2)
    add_body(doc, "The original ONNX model and generated C++ program will receive the same test input. Their outputs will be compared element by element using a documented absolute or combined tolerance. The report will include the maximum difference, class prediction agreement where applicable, and the final status: output preserved or validation failed.")

    add_heading(doc, "10 System Architecture", level=1)
    add_body(doc, "The architecture is organized around a source-format boundary and a target-code boundary. ONNX-specific information is consumed by the loader; the IR and later stages use ModelForge-owned structures.")
    add_table(
        doc,
        ["Stage", "Input", "Output"],
        [
            ("ONNX Loader", ".onnx model", "Graph, tensors, initializers, attributes"),
            ("Semantic Validator", "Loaded graph", "Validated graph or compiler diagnostic"),
            ("IR Builder", "Validated graph", "ModelForge IR instructions and tensor table"),
            ("Optimization Engine", "ModelForge IR", "Optimized IR and optimization statistics"),
            ("C++ Code Generator", "Optimized IR and constants", "model.cpp and model.h"),
            ("Build and Run", "Generated C++", "Standalone executable and prediction"),
            ("Validation Engine", "Original and generated outputs", "Difference report and pass/fail status"),
        ],
        widths=[1.65, 2.2, 2.9],
        font_size=8.7,
    )
    add_heading(doc, "10.1 Data flow", level=2)
    add_code_block(doc, [
        "ONNX model -> loader -> validated graph -> ModelForge IR",
        "             -> optimizer -> generated C++ -> native executable",
        "             -> original ONNX inference + generated inference",
        "             -> numerical comparison and validation report",
    ])
    add_heading(doc, "10.2 Module responsibilities", level=2)
    add_table(
        doc,
        ["Module", "Responsibility"],
        [
            ("ONNX Loader", "Read model metadata, graph nodes, tensors, weights, biases, and attributes."),
            ("Semantic Validator", "Check supported operators, tensor types, dimensions, required parameters, and graph consistency."),
            ("ModelForge IR", "Provide source-format-independent instructions and tensor metadata."),
            ("Optimizer", "Apply and count constant folding, redundant-operation removal, identity removal, and optional fusion."),
            ("Code Generator", "Emit weights, buffers, activation helpers, and the inference function."),
            ("Validation Engine", "Run both implementations, compare outputs, and report the numerical difference."),
            ("Visualization", "Optionally display the graph or a before-and-after IR view for explanation."),
        ],
        widths=[1.65, 5.1],
        font_size=8.8,
    )

    add_heading(doc, "11 Technology Stack", level=1)
    add_table(
        doc,
        ["Component", "Selected technology", "Reason"],
        [
            ("Compiler implementation", "C++17", "Matches the project brief and supports modular systems programming."),
            ("Source model format", "ONNX", "Portable graph representation for supported neural networks."),
            ("Serialization", "Protocol Buffers / ONNX protobuf", "Read the serialized model graph and initializers."),
            ("Build system", "CMake", "Portable configuration for the compiler and generated target."),
            ("Reference inference", "ONNX Runtime", "Baseline output for numerical equivalence checks."),
            ("Visualization", "Textual view or Graphviz", "Optional explanation of model and IR structure."),
            ("Version control", "Git and GitHub where permitted", "Record progress and preserve milestones."),
        ],
        widths=[1.75, 2.2, 2.8],
        font_size=8.8,
    )

    add_heading(doc, "12 Initial Prototype", level=1)
    add_body(doc, "The Phase 1 prototype is a small C++17 program stored in review_1_proposal_design/prototype/. It is intentionally dependency-free so the first demonstration can focus on compiler boundaries before the ONNX protobuf dependency is integrated.")
    add_table(
        doc,
        ["Prototype capability", "Demonstrated behavior"],
        [
            ("Model input", "Reads a small ModelForge manifest containing model, input, output, and node declarations."),
            ("Operator validation", "Accepts Gemm, MatMul, Add, Relu, Sigmoid, and Softmax; rejects unsupported operators."),
            ("Graph consistency", "Checks tensor production order, duplicate tensor names, and declared output availability."),
            ("IR boundary", "Prints INPUT, operator instructions, and RETURN in a simplified ModelForge IR view."),
            ("Error handling", "Prints a compiler-style error and returns a non-zero status for invalid input."),
        ],
        widths=[1.8, 4.95],
        font_size=9.0,
    )
    add_heading(doc, "12.1 Prototype input", level=2)
    add_code_block(doc, [
        "model iris_classifier",
        "input x float32 1,4",
        "output y float32 1,3",
        "node dense_1 Gemm x t1",
        "node activation_1 Relu t1 t2",
        "node dense_2 Gemm t2 t3",
        "node probabilities Softmax t3 y",
    ])
    doc.add_page_break()
    add_heading(doc, "12.2 Expected demonstration", level=2)
    add_code_block(doc, [
        "MODEL FORGE PHASE 1 PROTOTYPE",
        "Model: iris_classifier",
        "SUPPORTED OPERATORS",
        "  Gemm ................ OK",
        "  Relu ................ OK",
        "  Gemm ................ OK",
        "  Softmax ............. OK",
        "MODEL FORGE IR",
        "  INPUT  x",
        "  Gemm   input=x  output=t1",
        "  Relu   input=t1 output=t2",
        "  Gemm   input=t2 output=t3",
        "  Softmax input=t3 output=y",
        "  RETURN y",
        "STATUS: PROTOTYPE VALIDATION PASSED",
    ])
    add_body(doc, "The manifest is a Phase 1 bridge, not the final input format. In Phase 2, the loader will read ONNX protobuf data and feed the same validation and IR boundary. This separation makes the prototype useful while keeping the final architecture aligned with the stated ONNX-to-C++ objective.")

    add_heading(doc, "13 Feasibility and Innovation", level=1)
    add_heading(doc, "13.1 Technical feasibility", level=2)
    add_body(doc, "The first target model is small, static, and composed of well-understood operators. Dense layers, element-wise activations, and Softmax can be represented using ordinary C++ loops and arrays. The operator subset is small enough for explicit validation, and the generated code can avoid a full runtime by embedding the model parameters and inference logic.")
    add_heading(doc, "13.2 Risks and controls", level=2)
    add_table(
        doc,
        ["Risk", "Control"],
        [
            ("ONNX operator attributes differ across model exporters", "Freeze the initial operator subset and validate supported attributes explicitly."),
            ("Tensor shapes are not static or are incomplete", "Reject unsupported dynamic shapes with a precise diagnostic."),
            ("Floating-point outputs differ slightly", "Use a documented tolerance and report the maximum absolute difference."),
            ("Generated code is difficult to debug", "Prefer readable names, separate helpers, and deterministic formatting."),
            ("Scope expands beyond an individual project", "Protect the must-have pipeline and defer CNNs, GPU code, and full ONNX coverage."),
        ],
        widths=[2.55, 4.2],
        font_size=8.8,
    )
    add_heading(doc, "13.3 Individual contribution", level=2)
    add_body(doc, "The individual contribution is the design and implementation of a transparent educational ML compiler: a custom IR, ML-specific semantic diagnostics, explicit optimization reporting, readable C++ code generation, and automatic output-equivalence checking. The project is not presented as a replacement for industrial compilers or full ONNX runtimes.")

    add_heading(doc, "14 Review Roadmap", level=1)
    add_table(
        doc,
        ["Review", "Folder", "Evidence to prepare"],
        [
            ("Review 1", "review_1_proposal_design/", "Proposal, architecture, scope, methodology, prototype, feasibility, and viva preparation."),
            ("Review 2", "review_2_core_implementation/", "Working loader, validator, IR, optimizer, code generator, tests, and intermediate results."),
            ("Review 3", "review_3_final_integration/", "Complete compiler, generated code, validation logs, screenshots, benchmarks, report, and presentation."),
        ],
        widths=[1.0, 2.25, 3.5],
        font_size=8.9,
    )
    add_body(doc, "The project will be developed in the same order as the laboratory manual: define and design first, implement the core modules second, and integrate and test the complete system third. Each review folder is kept separate so the submitted evidence remains easy to locate.")

    add_heading(doc, "15 Review 1 Rubric Alignment", level=1)
    add_table(
        doc,
        ["Criterion", "Marks", "Evidence in this submission"],
        [
            ("Topic relevance", "3", "Compiler pipeline mapped to ONNX graph translation; scope tied to Compiler Design."),
            ("Problem statement", "3", "Clear problem, motivation, and core question in Sections 3 and 4."),
            ("Objectives", "2", "Measurable objectives listed in Section 5."),
            ("System design", "4", "Staged methodology, architecture, data flow, and module responsibilities in Sections 9 and 10."),
            ("Compiler concepts", "3", "Semantic analysis, custom IR, optimization, code generation, execution, and verification in Section 8."),
            ("Innovation", "2", "Educational ONNX-to-C++ compiler with transparent IR and output-preservation reporting."),
            ("Prototype", "3", "Dependency-free C++17 prototype with manifest validation and IR output in Section 12."),
        ],
        widths=[1.45, 0.65, 4.65],
        font_size=8.6,
    )
    add_callout(doc, "VIVA FOCUS", "Be ready to explain why the project uses a custom IR, why the supported operator set is intentionally small, how a shape mismatch becomes a compiler error, and how generated C++ output will be compared with ONNX Runtime output.")

    add_heading(doc, "16 Expected Outcomes", level=1)
    for item in [
        "A supported ONNX model can be loaded and inspected.",
        "Invalid operators, tensor types, dimensions, and graph references produce understandable diagnostics.",
        "The model is represented in a format independent of ONNX serialization.",
        "Optimization passes produce a before-and-after report without compromising correctness.",
        "Generated C++ source can be built as an independent inference program.",
        "Original and generated outputs agree within the selected numerical tolerance.",
        "The final demonstration clearly exposes the compiler stages and the student's individual understanding.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "17 Conclusion", level=1)
    add_body(doc, "ModelForge is a feasible and appropriately scoped Compiler Design Laboratory project. It treats an ONNX neural-network graph as a source representation, applies semantic analysis and IR-based transformations, and generates standalone C++ inference code. The first review establishes a clear architecture, a controlled operator subset, an implementation roadmap, and a working prototype boundary. The later reviews will add the ONNX loader, core compiler modules, generated code, and numerical validation evidence.")

    add_heading(doc, "References", level=1)
    refs = [
        "[1] Open Neural Network Exchange, Introduction to ONNX and ONNX Concepts, official documentation, https://onnx.ai/onnx/intro/ and https://onnx.ai/onnx/intro/concepts.html.",
        "[2] ONNX Runtime, official documentation and inference overview, https://onnxruntime.ai/docs/.",
        "[3] CMake, official reference documentation, https://cmake.org/cmake/help/latest/.",
        "[4] Compiler Design Laboratory Project Instruction Manual, supplied course document dated 03 September 2026.",
        "[5] ModelForge project brief, supplied file ModelForge_Project.md.",
    ]
    for ref in refs:
        add_body(doc, ref)

    add_heading(doc, "Appendix A Prototype File Map", level=1)
    add_code_block(doc, [
        "review_1_proposal_design/",
        "|-- Review_1_Project_Proposal_and_Design.docx",
        "`-- prototype/",
        "    |-- CMakeLists.txt",
        "    |-- prototype.cpp",
        "    |-- sample_model.mforge",
        "    `-- README.md",
    ])
    add_body(doc, "The shared project notes and the reserved Review 2 and Review 3 folders are located one level above this folder in the ModelForge project directory.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
