from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "review_1_proposal_design" / "Review_1_Project_Proposal_and_Design.docx"
FLOWCHART = ROOT / "review_1_proposal_design" / "model_forge_flowchart.png"

BLACK = "000000"
BODY = "000000"
GRAY = "000000"
LIGHT_GRAY = "F2F2F2"
GRID = "B7B7B7"


def set_font(run, name="Times New Roman", size=11, color=BODY, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def remove_paragraph_borders(style):
    ppr = style._element.get_or_add_pPr()
    for child in list(ppr):
        if child.tag == qn("w:pBdr"):
            ppr.remove(child)


def set_cell_border(cell, color=GRID, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn("w:" + edge)
        element = borders.find(tag)
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_text(cell, text, bold=False, size=10.2, color=BODY, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_font(r, size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(doc, headers, rows, widths, size=10.0):
    """Render tabular information as clean labeled paragraphs without borders."""
    for values in rows:
        if len(headers) == 2 and headers == ["Initial scope", "Outside the initial scope"]:
            for label, value in zip(headers, values):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.25)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(f"{label}: ")
                set_font(r, size=size, color=BLACK, bold=True)
                r = p.add_run(value)
                set_font(r, size=size, color=BLACK)
        elif len(headers) == 2:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.25)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(f"{values[0]}: ")
            set_font(r, size=size, color=BLACK, bold=True)
            r = p.add_run(values[1])
            set_font(r, size=size, color=BLACK)
        elif len(headers) == 3:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.25)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(values[0])
            set_font(r, size=size, color=BLACK, bold=True)
            r = p.add_run(f" ({values[1]}): ")
            set_font(r, size=size, color=BLACK)
            r = p.add_run(values[2])
            set_font(r, size=size, color=BLACK)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.25)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(" — ".join(values))
            set_font(r, size=size, color=BLACK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_code(doc, lines):
    for index, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_together = True
        p.paragraph_format.keep_with_next = index < len(lines) - 1
        r = p.add_run(line)
        set_font(r, name="Consolas", size=8.8, color=BLACK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    set_font(run, size=9, color=GRAY)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after in (("Title", 22, 0, 8), ("Heading 1", 15, 14, 5), ("Heading 2", 12, 9, 4), ("Heading 3", 11, 7, 3)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        remove_paragraph_borders(style)

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(BODY)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.1


def configure_page(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    hp = section.header.paragraphs[0]
    hp.text = ""

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    r = fp.add_run("Compiler Design Laboratory  |  ")
    set_font(r, size=9, color=GRAY)
    add_page_number(fp)


def add_body(doc, text, style="Body Text"):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.widow_control = True
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.widow_control = True
    p.add_run(text)
    return p


def add_numbered(doc, items):
    for index, text in enumerate(items, start=1):
        p = doc.add_paragraph(style="Body Text")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.widow_control = True
        p.add_run(f"{index}. ")
        p.add_run(text)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.widow_control = True
    p.add_run(text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_identity_block(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("ModelForge")
    set_font(r, name="Times New Roman", size=22, color=BLACK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Compiler Design Phase 1")
    set_font(r, name="Times New Roman", size=14, color=BLACK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Name: D.N.S.Vaishnavi")
    set_font(r, name="Times New Roman", size=11, color=BLACK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("Reg no: 24BCB0093")
    set_font(r, name="Times New Roman", size=11, color=BLACK)


def build():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    doc.core_properties.title = "ModelForge Review 1 Project Proposal and Design"
    doc.core_properties.subject = "Compiler Design Laboratory project proposal"
    doc.core_properties.author = ""

    add_identity_block(doc)

    add_heading(doc, "Abstract")
    add_body(doc, "ModelForge is an educational compiler for a small, controlled subset of ONNX feed-forward neural-network models. It will read a model graph, validate supported operators and tensor properties, lower the graph into a custom intermediate representation, apply simple optimizations, and generate readable standalone C++17 inference code.")
    add_body(doc, "The generated program will be compared with the original model executed through ONNX Runtime. The purpose is not to reproduce a complete industrial ML compiler. The purpose is to make the stages of model loading, semantic analysis, intermediate representation, optimization, target-code generation, and output verification visible in one manageable project.")

    add_heading(doc, "1. Problem Statement")
    add_body(doc, "Small machine-learning models are usually executed through general-purpose inference runtimes. This is practical, but it hides the translation from a model graph to executable computations and introduces dependencies that are unnecessary for a small demonstration network. A compact compiler is needed to show how a model can be inspected, validated, transformed, and translated into understandable target code.")
    add_body(doc, "ModelForge addresses the following question: how can a small ONNX neural-network model be translated into standalone C++ inference code while demonstrating compiler-design concepts such as semantic validation, intermediate representation, optimization, target-code generation, and correctness verification?")

    add_heading(doc, "2. Motivation")
    add_body(doc, "Compiler Design projects often begin with arithmetic expressions or a small programming language. ModelForge applies the same ideas to a machine-learning graph. A graph node consumes tensors, produces tensors, and carries attributes or constants; these relationships must be checked before the computation is translated. This makes the project a practical example of compiler construction applied to a current software format.")
    add_body(doc, "The project is deliberately limited to small feed-forward networks. That limitation keeps the implementation understandable for an individual project and makes the final demonstration easier to explain during the viva. Readable generated code and clear diagnostics are more important than broad operator coverage.")

    add_heading(doc, "3. Objectives")
    add_numbered(doc, [
        "Read a supported ONNX neural-network model.",
        "Extract graph inputs, outputs, operators, initializers, weights, biases, types, shapes, and relevant attributes.",
        "Detect unsupported operators and invalid tensor properties with compiler-style diagnostics.",
        "Lower the model into a simplified ModelForge intermediate representation.",
        "Implement at least two understandable optimization passes and report their effects.",
        "Generate readable standalone C++17 source and header files for inference.",
        "Compile and execute the generated program during the integrated project.",
        "Compare generated output with the original ONNX model within a documented tolerance.",
    ])

    add_heading(doc, "4. Scope")
    add_body(doc, "The first demonstration model is an Iris-style classifier with four input values, a dense layer of eight neurons, a ReLU activation, a dense layer of three neurons, and a three-class Softmax output.")
    add_table(
        doc,
        ["Initial scope", "Outside the initial scope"],
        [
            ("Gemm, MatMul, Add, Relu, Sigmoid, and Softmax", "CNN-heavy architectures such as YOLO and ResNet"),
            ("Static float32 shapes", "Transformers, attention, LSTM, and GRU"),
            ("Small feed-forward classifiers", "Dynamic control flow and arbitrary dynamic shapes"),
            ("C++17 CPU code generation", "GPU or CUDA code generation"),
            ("Textual or simple Graphviz visualization", "Complete ONNX operator coverage"),
        ],
        widths=[3.25, 3.3],
        size=9.6,
    )
    add_body(doc, "Possible future additions include Flatten, Tanh, Identity, an MNIST MLP, memory reuse, SIMD generation, and richer visualization. These will be considered only after the core pipeline is working.")

    add_heading(doc, "5. Background Study")
    add_heading(doc, "5.1 ONNX model representation", level=2)
    add_body(doc, "ONNX represents a model using graph inputs and outputs, operator nodes, initializers, attributes, and tensor types. The graph is serialized using Protocol Buffers and uses versioned operator semantics. This structure makes ONNX a suitable source representation for a teaching compiler: the loader can translate the graph into compiler-owned structures while later stages remain independent of the file format. [1]")
    add_heading(doc, "5.2 Reference inference", level=2)
    add_body(doc, "ONNX Runtime provides a standard execution path for ONNX models and has APIs for several languages, including C and C++. ModelForge will use the original model executed by ONNX Runtime as the reference when checking the output of generated C++ code. It is a validation baseline, not the target generated by the compiler. [2]")
    add_heading(doc, "5.3 Build system", level=2)
    add_body(doc, "CMake will be used to configure the compiler and the generated target program. Keeping the compiler executable and generated model as separate build targets will make the final demonstration show both stages clearly. [3]")

    add_heading(doc, "6. Compiler Design Concepts")
    add_table(
        doc,
        ["Concept", "Use in ModelForge"],
        [
            ("Model loading", "Read the serialized graph and convert it into internal C++ structures."),
            ("Semantic analysis", "Check supported operators, data types, dimensions, required tensors, and attributes."),
            ("Intermediate representation", "Represent inputs, constants, dense operations, activations, Softmax, and returns independently of ONNX."),
            ("Optimization", "Apply constant folding, redundant-operation removal, identity removal, and selected operation fusion."),
            ("Target-code generation", "Emit readable C++17 functions, buffers, constants, and activation helpers."),
            ("Execution", "Build the generated source and run standalone inference."),
            ("Verification", "Compare original and generated outputs using a documented floating-point tolerance."),
            ("Error handling", "Report unsupported operators, missing tensors, shape mismatches, and invalid attributes."),
        ],
        widths=[2.0, 4.55],
        size=9.4,
    )
    add_body(doc, "Lexical analysis and grammar parsing are not central here because ONNX is a serialized graph rather than a text language written by the user. The project instead focuses on the compiler stages that fit graph translation most directly: semantic analysis, IR construction, optimization, and target generation.")

    add_heading(doc, "7. Proposed Methodology")
    add_body(doc, "The compiler will use the following sequence. Each stage passes a defined structure to the next stage so that ONNX-specific logic does not spread into the optimizer or code generator.")
    add_numbered(doc, [
        "Load the ONNX model and collect graph metadata, nodes, tensors, initializers, and attributes.",
        "Validate operators, data types, dimensions, tensor references, and supported attributes.",
        "Lower the validated graph into ModelForge IR instructions and a tensor table.",
        "Run optimization passes on the IR and record before-and-after instruction counts.",
        "Generate model.cpp and model.h containing constants, buffers, helper functions, and inference logic.",
        "Compile and run the generated program, then compare its output with the original ONNX model.",
    ])
    add_heading(doc, "7.1 Validation strategy", level=2)
    add_body(doc, "The validator will reject operators outside the initial supported set, incompatible matrix dimensions, unsupported data types, missing initializers, duplicate tensor names, missing model outputs, and invalid Softmax axes. A diagnostic will name the failing node or tensor and state the expected condition.")
    add_heading(doc, "7.2 Optimization strategy", level=2)
    add_body(doc, "The first passes will be constant folding, redundant-operation elimination, and identity removal. A later pass may fuse a dense operation with a following ReLU when the graph pattern and tensor use are safe. The optimization report will separate an IR simplification from a measured runtime improvement.")
    add_heading(doc, "7.3 Correctness strategy", level=2)
    add_body(doc, "The same input will be passed to the original ONNX model and the generated C++ program. Outputs will be compared element by element, and the report will include the maximum difference, class prediction agreement where applicable, the tolerance, and the final pass or fail status.")

    add_heading(doc, "8. System Architecture")
    architecture_intro = add_body(doc, "The architecture has a clear source-format boundary. ONNX-specific parsing stops at the loader; the IR, optimizer, and code generator use ModelForge structures. Visualization is optional and reads the graph or IR for explanation only.")
    architecture_intro.paragraph_format.keep_with_next = True
    flowchart_paragraph = doc.add_paragraph()
    flowchart_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    flowchart_paragraph.paragraph_format.keep_together = True
    flowchart_paragraph.paragraph_format.keep_with_next = True
    flowchart_run = flowchart_paragraph.add_run()
    flowchart_run.add_picture(str(FLOWCHART), width=Inches(6.45))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(6)
    caption_run = caption.add_run("Figure 1. ModelForge compiler flowchart.")
    set_font(caption_run, size=9.5, color=BLACK, italic=True)
    add_table(
        doc,
        ["Module", "Responsibility"],
        [
            ("ONNX Loader", "Read model metadata, graph nodes, tensors, weights, biases, and attributes."),
            ("Semantic Validator", "Check supported operators, tensor types, dimensions, required parameters, and graph consistency."),
            ("ModelForge IR", "Provide source-format-independent instructions and tensor metadata."),
            ("Optimizer", "Apply and count constant folding, redundant-operation removal, identity removal, and optional fusion."),
            ("C++ Code Generator", "Emit weights, buffers, activation helpers, and the inference function."),
            ("Validation Engine", "Run both implementations, compare outputs, and report the numerical difference."),
            ("Visualization", "Optionally display the graph or a before-and-after IR view."),
        ],
        widths=[2.0, 4.55],
        size=9.5,
    )

    add_heading(doc, "9. Technology Stack")
    add_table(
        doc,
        ["Component", "Technology", "Purpose"],
        [
            ("Compiler", "C++17", "Main implementation language."),
            ("Model format", "ONNX", "Portable source representation for supported models."),
            ("Serialization", "Protocol Buffers / ONNX protobuf", "Read the serialized graph and initializers."),
            ("Build", "CMake", "Configure the compiler and generated target."),
            ("Reference inference", "ONNX Runtime", "Provide baseline output for equivalence checks."),
            ("Visualization", "Text output or Graphviz", "Support explanation of graph and IR structure."),
            ("Version control", "Git", "Record milestones and preserve project history."),
        ],
        widths=[1.75, 2.25, 2.55],
        size=9.4,
    )

    add_heading(doc, "10. Conclusion")
    add_body(doc, "ModelForge is a feasible and appropriately scoped Compiler Design Laboratory project. It treats an ONNX neural-network graph as a source representation, applies semantic analysis and IR-based transformations, and generates standalone C++ inference code.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
